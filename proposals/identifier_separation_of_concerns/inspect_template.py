from pptx import Presentation
from pptx.util import Inches, Pt, Emu
import json

prs = Presentation(r"C:\Users\aluk\Downloads\AOUSD 2025_PowerPoint Template.pptx")

print(f"Slide width: {prs.slide_width}, height: {prs.slide_height}")
print(f"Number of slide layouts: {len(prs.slide_layouts)}")
print(f"Number of slides: {len(prs.slides)}")
print()

for i, layout in enumerate(prs.slide_layouts):
    print(f"Layout {i}: '{layout.name}'")
    for ph in layout.placeholders:
        print(f"  Placeholder {ph.placeholder_format.idx}: '{ph.name}' type={ph.placeholder_format.type} "
              f"pos=({ph.left},{ph.top}) size=({ph.width},{ph.height})")
    print()

print("=" * 60)
print("SLIDES:")
print("=" * 60)
for i, slide in enumerate(prs.slides):
    print(f"\nSlide {i}: layout='{slide.slide_layout.name}'")
    for shape in slide.shapes:
        print(f"  Shape: '{shape.name}' type={shape.shape_type} "
              f"pos=({shape.left},{shape.top}) size=({shape.width},{shape.height})")
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                text = p.text.strip()
                if text:
                    font_info = ""
                    if p.runs:
                        r = p.runs[0]
                        font_info = f" [font={r.font.name}, size={r.font.size}, bold={r.font.bold}]"
                    print(f"    '{text[:80]}'{font_info}")
