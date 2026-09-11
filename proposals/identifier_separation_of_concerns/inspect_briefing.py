from pptx import Presentation
from pptx.util import Emu
import re

prs = Presentation(r"C:\Users\aluk\Downloads\2026 - AOUSD Overview Briefing Deck (1).pptx")

print(f"Slide size: {Emu(prs.slide_width).inches:.2f} x {Emu(prs.slide_height).inches:.2f} in")
print(f"Slides: {len(prs.slides)}")
print(f"Layouts: {len(prs.slide_layouts)}")
print()

for i, layout in enumerate(prs.slide_layouts):
    print(f"Layout {i}: '{layout.name}'")

print()
print("=" * 70)

for i, slide in enumerate(prs.slides):
    layout_name = slide.slide_layout.name
    print(f"\n{'='*70}")
    print(f"SLIDE {i}: layout='{layout_name}'")
    print(f"{'='*70}")
    for shape in slide.shapes:
        stype = str(shape.shape_type) if hasattr(shape, 'shape_type') else '?'
        pos = f"({Emu(shape.left).inches:.1f}, {Emu(shape.top).inches:.1f})"
        sz = f"{Emu(shape.width).inches:.1f}x{Emu(shape.height).inches:.1f}"
        print(f"  [{stype}] '{shape.name}' at {pos} size={sz}")

        if hasattr(shape, 'fill'):
            try:
                fill = shape.fill
                if fill.type is not None:
                    print(f"    fill: type={fill.type}")
                    if hasattr(fill, 'fore_color') and fill.fore_color and fill.fore_color.type is not None:
                        try:
                            print(f"    fill color: #{fill.fore_color.rgb}")
                        except:
                            print(f"    fill color: theme={fill.fore_color.theme_color}")
            except:
                pass

        if hasattr(shape, 'line'):
            try:
                if shape.line.color and shape.line.color.type is not None:
                    try:
                        print(f"    line: #{shape.line.color.rgb} width={shape.line.width}")
                    except:
                        pass
            except:
                pass

        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                text = p.text.strip()
                if text:
                    font_info = ""
                    if p.runs:
                        r = p.runs[0]
                        parts = []
                        if r.font.name: parts.append(f"font={r.font.name}")
                        if r.font.size: parts.append(f"size={r.font.size}")
                        if r.font.bold: parts.append("bold")
                        if r.font.italic: parts.append("italic")
                        try:
                            if r.font.color and r.font.color.rgb:
                                parts.append(f"color=#{r.font.color.rgb}")
                        except:
                            pass
                        font_info = f" [{', '.join(parts)}]"
                    print(f"    '{text[:80]}'{font_info}")
