from pptx import Presentation
from pptx.enum.text import PP_ALIGN

prs = Presentation(r"C:\Users\aluk\Downloads\AOUSD 2025_PowerPoint Template.pptx")

# Check existing transition slides for alignment
for i, slide in enumerate(prs.slides):
    if slide.slide_layout.name == 'Transition':
        print(f"Slide {i}: layout='Transition'")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    print(f"  alignment={p.alignment}, text='{p.text[:60]}'")
                print(f"  word_wrap={shape.text_frame.word_wrap}")
                print(f"  auto_size={shape.text_frame.auto_size}")

# Also check the layout placeholder defaults
layout = None
for l in prs.slide_layouts:
    if l.name == 'Transition':
        layout = l
        break

print(f"\nLayout placeholder:")
for ph in layout.placeholders:
    print(f"  idx={ph.placeholder_format.idx} pos=({ph.left},{ph.top}) size=({ph.width},{ph.height})")
    if ph.has_text_frame:
        for p in ph.text_frame.paragraphs:
            print(f"    alignment={p.alignment}")
