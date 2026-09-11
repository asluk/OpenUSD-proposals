from pptx import Presentation
from pptx.util import Emu
from lxml import etree

prs = Presentation(r"C:\Users\aluk\Downloads\AOUSD 2025_PowerPoint Template.pptx")

w = prs.slide_width
h = prs.slide_height
print(f"Width: {w} EMU = {Emu(w).inches:.2f} in")
print(f"Height: {h} EMU = {Emu(h).inches:.2f} in")
print()

for i, layout in enumerate(prs.slide_layouts):
    print(f"Layout {i}: '{layout.name}'")

print()

ns = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
}

print("=== THEME COLORS ===")
theme = prs.slide_masters[0].slide_layouts[0].slide_master.element
for clr in theme.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}clrScheme'):
    for child in clr:
        tag = child.tag.split('}')[-1]
        for sub in child:
            val = sub.get('val', sub.get('lastClr', ''))
            print(f"  {tag}: #{val}")

print()
print("=== ACCENT COLORS FROM SHAPES ===")
for i, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        xml = shape._element.xml
        for tag in ['srgbClr', 'schemeClr']:
            import re
            matches = re.findall(rf'<a:{tag}\s+val="([^"]+)"', xml)
            for m in matches:
                if tag == 'srgbClr' and m not in ('000000', 'FFFFFF', 'ffffff'):
                    print(f"  Slide {i}, shape '{shape.name}': #{m}")
                elif tag == 'schemeClr':
                    pass
