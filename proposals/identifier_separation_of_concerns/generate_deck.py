from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

TEMPLATE = r"C:\Users\aluk\Downloads\AOUSD 2025_PowerPoint Template.pptx"
OUTPUT = r"c:\git\OpenUSD-proposals\proposals\identifier_separation_of_concerns\Identifier_Separation_of_Concerns.pptx"

BLUE_DARK = RGBColor(0x56, 0x5D, 0x95)
BLUE_LIGHT = RGBColor(0xBB, 0xD6, 0xEE)
BLUE_FILL = RGBColor(0xE8, 0xF0, 0xF8)
GRAY_TEXT = RGBColor(0x3F, 0x3F, 0x3F)
GRAY_BG = RGBColor(0xF5, 0xF5, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)

prs = Presentation(TEMPLATE)

SW = prs.slide_width   # 36576000 EMU = 40 in
SH = prs.slide_height  # 20574000 EMU = 22.5 in

sldIdLst = prs.slides._sldIdLst
slide_rIds = []
for sldId in list(sldIdLst):
    rId = sldId.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    slide_rIds.append(rId)
for rId in slide_rIds:
    for sldId in list(sldIdLst):
        rid = sldId.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        if rid == rId:
            sldIdLst.remove(sldId)
            break
    prs.part.drop_rel(rId)

layouts = {l.name: l for l in prs.slide_layouts}


# ── helper: accent bar under title ─────────────────────────────
def add_accent_bar(slide, top=Inches(3.8), width=Inches(4), left=Inches(2)):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.12)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE_DARK
    bar.line.fill.background()
    return bar


# ── helper: rounded-rect card ──────────────────────────────────
ICON_DIR = r"c:\git\OpenUSD-proposals\proposals\identifier_separation_of_concerns\icons"


FONT_TITLE = "Montserrat SemiBold"
FONT_BODY = "Montserrat"


def add_card(slide, left, top, width, height, heading, bullets,
             fill=WHITE, border=BLUE_LIGHT,
             head_size=Pt(80), bullet_size=Pt(56), icon=None,
             header_bar=True):
    import os

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = fill
    card.line.color.rgb = border
    card.line.width = Pt(2.5)
    adj = card.adjustments
    if len(adj) > 0:
        adj[0] = 0.02

    margin_x = Inches(0.6)
    margin_top = Inches(0.4)
    tb_width = width - 2 * margin_x

    bar_h = Inches(0.15)
    if header_bar:
        bar = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left + Inches(0.3), top + Inches(0.3),
            width - Inches(0.6), bar_h
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = BLUE_DARK
        bar.line.fill.background()
        if len(bar.adjustments) > 0:
            bar.adjustments[0] = 0.5

    cursor_y = top + margin_top + bar_h + Inches(0.3)

    if icon:
        icon_path = os.path.join(ICON_DIR, icon)
        if os.path.exists(icon_path):
            from PIL import Image as PILImage
            img = PILImage.open(icon_path)
            aspect = img.width / img.height
            icon_box = Inches(2.8)
            if aspect >= 1:
                icon_w = icon_box
                icon_h = int(icon_box / aspect)
            else:
                icon_h = icon_box
                icon_w = int(icon_box * aspect)
            icon_x = left + (width - icon_w) // 2
            icon_y = cursor_y + (icon_box - icon_h) // 2
            slide.shapes.add_picture(icon_path, icon_x, icon_y, icon_w, icon_h)
            cursor_y += icon_box + Inches(0.2)

    head_h = Inches(1.6)
    txBox = slide.shapes.add_textbox(left + margin_x, cursor_y, tb_width, head_h)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    p = tf.paragraphs[0]
    p.text = heading.upper()
    p.font.size = head_size
    p.font.bold = True
    p.font.name = FONT_TITLE
    p.font.color.rgb = BLACK
    p.alignment = PP_ALIGN.CENTER if icon else PP_ALIGN.LEFT
    cursor_y += head_h + Inches(0.1)

    bullet_h = height - (cursor_y - top) - Inches(0.4)
    bBox = slide.shapes.add_textbox(left + margin_x, cursor_y, tb_width, bullet_h)
    bf = bBox.text_frame
    bf.word_wrap = True
    for i, b in enumerate(bullets):
        p = bf.paragraphs[0] if i == 0 else bf.add_paragraph()
        p.text = b
        p.font.size = bullet_size
        p.font.name = FONT_BODY
        p.font.color.rgb = GRAY_TEXT
        p.space_after = Pt(14)

    return card


# ── helper: title + subtitle on blank slide ────────────────────
def add_headed_blank(title_text, subtitle_text=""):
    slide = prs.slides.add_slide(layouts['No BG Image'])

    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid()
    bg.fill.fore_color.rgb = GRAY_BG
    bg.line.fill.background()

    txBox = slide.shapes.add_textbox(
        Inches(1.5), Inches(0.8), Inches(37), Inches(2.5)
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(96)
    p.font.bold = True
    p.font.name = FONT_TITLE
    p.font.color.rgb = BLACK

    if subtitle_text:
        stBox = slide.shapes.add_textbox(
            Inches(1.5), Inches(3.3), Inches(37), Inches(1.5)
        )
        sf = stBox.text_frame
        sf.word_wrap = True
        p = sf.paragraphs[0]
        p.text = subtitle_text
        p.font.size = Pt(56)
        p.font.name = FONT_BODY
        p.font.color.rgb = GRAY_TEXT

    add_accent_bar(slide, top=Inches(4.8), width=Inches(5), left=Inches(1.5))
    return slide


# ── template layout helpers (unchanged) ────────────────────────
def add_title_slide(title, subtitle):
    slide = prs.slides.add_slide(layouts['Title Slide'])
    slide.placeholders[0].text = title
    slide.placeholders[1].text = subtitle
    return slide


def add_transition(text):
    slide = prs.slides.add_slide(layouts['Transition'])
    slide.placeholders[1].text = text
    slide.placeholders[1].text_frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    slide.placeholders[1].text_frame.word_wrap = True
    return slide


def add_slide(title, subtitle, bullets):
    slide = prs.slides.add_slide(layouts['Title, Subtitle, Bullets'])
    slide.placeholders[0].text = title
    slide.placeholders[1].text = subtitle
    tf = slide.placeholders[2].text_frame
    tf.clear()
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
    return slide


# ============================================================
# NARRATIVE ARC
# ============================================================

# 1 - TITLE
add_title_slide(
    "Separation of Concerns for Identifiers in USD",
    "Aaron Luk  \u2022  NVIDIA  \u2022  Alliance for OpenUSD"
)

# 2 - HOOK
add_transition("The Problem")

# 3 - THE TENSION
add_slide(
    "Every System Has Its Own Notion of Identity",
    "USD has one kind of identifier. The world has thousands.",
    [
        "A Revit ElementId. An IFC GlobalId. A PLM part number.",
        "An OPC UA NodeId. A shot-tracking database key.",
        "They all point at \u201cthe same thing\u201d \u2014 but they don\u2019t agree "
        "on what that thing is.",
        "",
        "USD namespace paths were never meant to carry all of this.",
        "And yet, today, there\u2019s nowhere else to put it.",
    ]
)

# 4 - NOT A BUG
add_slide(
    "This Is Not a Flaw in USD",
    "It was never specific to USD \u2014 and nothing before USD could solve it",
    [
        "Every scene graph and interchange format has struggled here",
        "USD is the first with the composition model, cross-industry "
        "adoption, and standards body to solve it properly",
        "The GUID debate? A symptom, not the disease.",
    ]
)

# 5 - AHA
add_transition("The Separation of Concerns")

# 6 - SEPARATION OF CONCERNS (card layout)
slide = add_headed_blank(
    "Two Problems Are Being Conflated",
    "Separating them changes everything"
)
card_top = Inches(6.0)
card_h = Inches(14.0)
card_w = Inches(17.0)
gap = Inches(2.0)
left_start = Inches(2.0)

add_card(slide, left_start, card_top, card_w, card_h,
         "USD Namespace Path",
         [
             "Unique per prim instance",
             "Governed by grammar rules (XID)",
             "Primary key for composition",
             "Tied to position in hierarchy",
         ],
         fill=BLUE_FILL, border=BLUE_DARK)

add_card(slide, left_start + card_w + gap, card_top, card_w, card_h,
         "Source Identifier",
         [
             "May be shared across instances",
             "Any characters, any format",
             "Traceability, BOM, compliance",
             "Survives renames and refactors",
         ],
         fill=WHITE, border=BLUE_DARK)

# 7 - WHY NOW
add_slide(
    "Why Now",
    "The window is open \u2014 and closing",
    [
        "displayName is migrating into uiHints (UI Hints, implemented)",
        "Pipelines using it for source IDs face API breakage today",
        "Every month without a standard = more ad-hoc debt",
        "New industries are adopting USD faster than conventions can form",
    ]
)

# 8 - TRANSITION
add_transition("Cross-Industry Evidence")

# 9 - USE CASES (clean icon grid, briefing-deck style)
slide = prs.slides.add_slide(layouts['Title, Subtitle, Bullets'])
slide.placeholders[0].text = "The Problem Shows Up Everywhere"
slide.placeholders[1].text = "Different domains, same gap \u2014 the common thread is traceability"
slide.placeholders[2].text_frame.clear()
slide.placeholders[2].text_frame.paragraphs[0].text = ""

import os
from PIL import Image as PILImage

domains = [
    ("AECO", [
        "IFC GlobalIds",
        "Classification codes",
        "Revision workflows",
    ], "m48-data-center.png"),
    ("Manufacturing", [
        "Part & serial numbers",
        "BOM traceability",
        "Identity across processes",
    ], "m48-robot-manufacturing.png"),
    ("Digital\nEngineering", [
        "Equipment tags",
        "OPC UA NodeIds",
        "Telemetry bindings",
    ], "m48-workflow.png"),
    ("Media &\nEntertainment", [
        "Asset database IDs",
        "Shot & sequence tracking",
        "Production management",
    ], "m48-render.png"),
]

n = len(domains)
col_w = Inches(9.0)
col_gap = Inches(0.5)
total_w = n * col_w + (n - 1) * col_gap
grid_left = (SW - total_w) // 2
icon_row_y = Inches(6.5)
icon_box = Inches(3.0)
label_y = icon_row_y + icon_box + Inches(0.3)
detail_y = label_y + Inches(2.0)

for idx, (label, details, icon_file) in enumerate(domains):
    cx = grid_left + idx * (col_w + col_gap) + col_w // 2

    icon_path = os.path.join(ICON_DIR, icon_file)
    if os.path.exists(icon_path):
        img = PILImage.open(icon_path)
        aspect = img.width / img.height
        if aspect >= 1:
            iw = icon_box
            ih = int(icon_box / aspect)
        else:
            ih = icon_box
            iw = int(icon_box * aspect)
        slide.shapes.add_picture(
            icon_path,
            cx - iw // 2, icon_row_y + (icon_box - ih) // 2,
            iw, ih
        )

    tb = slide.shapes.add_textbox(cx - col_w // 2, label_y, col_w, Inches(1.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = label
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.name = FONT_TITLE
    p.font.color.rgb = BLACK
    p.alignment = PP_ALIGN.CENTER

    dt = slide.shapes.add_textbox(cx - col_w // 2, detail_y, col_w, Inches(3.0))
    df = dt.text_frame
    df.word_wrap = True
    for i, line in enumerate(details):
        p = df.paragraphs[0] if i == 0 else df.add_paragraph()
        p.text = line
        p.font.size = Pt(30)
        p.font.name = FONT_BODY
        p.font.color.rgb = GRAY_TEXT
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(6)

# 10 - KEY INSIGHT
add_slide(
    "Source Identity Is the Primary Gap",
    "Instance identity already works. Source identity doesn\u2019t exist yet.",
    [
        "30 doors, same catalog number \u2014 unique prim paths, shared source identity",
        "Real objects carry identifiers from multiple systems simultaneously",
        "A chiller unit: PLM part number + IFC GlobalId + OPC UA NodeId",
        "This needs a mechanism keyed by domain, not a single field",
    ]
)

# 11 - TRANSITION
add_transition("Candidate Approaches")

# 12 - APPROACHES (two cards)
slide = add_headed_blank(
    "Approach A vs. Approach B",
    "Evaluating both before committing"
)

app_top = Inches(6.0)
app_h = Inches(14.5)
app_w = Inches(17.0)
app_gap = Inches(2.0)
app_left = Inches(2.0)

add_card(slide, app_left, app_top, app_w, app_h,
         "Extend assetInfo",
         [
             "Stratified sub-dictionaries",
             "",
             "\u2713  Existing composition semantics",
             "\u2713  Low barrier to adoption",
             "\u2713  Domains move independently",
             "",
             "\u2717  Not in UsdPrimDefinition",
             "\u2717  Dumping-ground risk (more acute)",
         ],
         fill=BLUE_FILL, border=BLUE_DARK)

add_card(slide, app_left + app_w + app_gap, app_top, app_w, app_h,
         "Leverage Applied Schemas",
         [
             "Typed properties",
             "",
             "\u2713  Full discoverability",
             "\u2713  Schema versioning & validation",
             "\u2713  More rigorous structure",
             "",
             "\u2717  Agreement on schema slows adoption",
             "\u2717  Heterogeneous packages are hard",
             "\u2717  Requires distributing plugins",
         ],
         fill=WHITE, border=BLUE_DARK)

# 13 - THE HARD PART
add_slide(
    "The Hard Part Isn\u2019t Technical",
    "Both approaches require curation \u2014 and curation slows production",
    [
        "No one wants to wait for consensus before proving what works",
        "The design must let domains move independently\u2026",
        "\u2026while still converging on interoperable conventions over time",
        "",
        "AI can help derive linkages across uncurated identifiers \u2014",
        "but only human agreement makes the answer verifiable",
    ]
)

# 14 - DESIGN PRINCIPLES (card grid: 4 + 3)
slide = add_headed_blank(
    "Design Principles",
    "Seven guardrails for the evaluation"
)

principles = [
    "Separation\nof concerns",
    "Industry\nagnosticism",
    "Composability",
    "Discoverability",
    "External\nqueryability",
    "Round-trip\nfidelity",
    "Minimal\ndisruption",
]

pr_w = Inches(8.0)
pr_h = Inches(6.0)
pr_gap = Inches(0.8)
row1_y = Inches(6.0)
row2_y = row1_y + pr_h + pr_gap

total_4 = 4 * pr_w + 3 * pr_gap
row1_left = (SW - total_4) // 2
total_3 = 3 * pr_w + 2 * pr_gap
row2_left = (SW - total_3) // 2

for i in range(4):
    x = row1_left + i * (pr_w + pr_gap)
    num = str(i + 1)
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, x, row1_y, pr_w, pr_h
    )
    card.fill.solid()
    card.fill.fore_color.rgb = BLUE_FILL if i % 2 == 0 else WHITE
    card.line.color.rgb = BLUE_LIGHT
    card.line.width = Pt(2)
    if len(card.adjustments) > 0:
        card.adjustments[0] = 0.05
    tb = slide.shapes.add_textbox(x, row1_y + Inches(0.4), pr_w, Inches(1.2))
    p = tb.text_frame.paragraphs[0]
    p.text = num
    p.font.size = Pt(72)
    p.font.bold = True
    p.font.color.rgb = BLUE_DARK
    p.alignment = PP_ALIGN.CENTER
    tb2 = slide.shapes.add_textbox(x, row1_y + Inches(1.8), pr_w, Inches(3.8))
    p2 = tb2.text_frame.paragraphs[0]
    p2.text = principles[i]
    p2.font.size = Pt(48)
    p2.font.color.rgb = GRAY_TEXT
    p2.alignment = PP_ALIGN.CENTER
    tb2.text_frame.word_wrap = True

for i in range(3):
    x = row2_left + i * (pr_w + pr_gap)
    num = str(i + 5)
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, x, row2_y, pr_w, pr_h
    )
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE if i % 2 == 0 else BLUE_FILL
    card.line.color.rgb = BLUE_LIGHT
    card.line.width = Pt(2)
    if len(card.adjustments) > 0:
        card.adjustments[0] = 0.05
    tb = slide.shapes.add_textbox(x, row2_y + Inches(0.4), pr_w, Inches(1.2))
    p = tb.text_frame.paragraphs[0]
    p.text = num
    p.font.size = Pt(72)
    p.font.bold = True
    p.font.color.rgb = BLUE_DARK
    p.alignment = PP_ALIGN.CENTER
    tb2 = slide.shapes.add_textbox(x, row2_y + Inches(1.8), pr_w, Inches(3.8))
    p2 = tb2.text_frame.paragraphs[0]
    p2.text = principles[i + 4]
    p2.font.size = Pt(48)
    p2.font.color.rgb = GRAY_TEXT
    p2.alignment = PP_ALIGN.CENTER
    tb2.text_frame.word_wrap = True

# 15 - TRANSITION
add_transition("What\u2019s Next")

# 16 - NEXT STEPS
add_slide(
    "Next Steps",
    "Consensus at every step \u2014 stakeholder engagement accelerates the work",
    [
        "1. Submit as PR to OpenUSD-proposals",
        "2. Align on the problem statement across industries",
        "3. Gather additional use cases",
        "4. Evaluate Approach A vs. B (or hybrid)",
        "5. Draft a solution proposal with mechanism, semantics, and API",
    ]
)

# 17 - CLOSING
prs.slides.add_slide(layouts['Closing Slide'])

prs.save(OUTPUT)
print(f"Saved to {OUTPUT}")
print(f"Total slides: {len(prs.slides)}")
